"""Excel ingestion: the real workbook and synthetic layouts."""

from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl
import pytest

from src.config import PROJECT_ROOT, REQUIRED_WORKFLOW_IDS
from src.errors import CatalogError
from src.excel_loader import load_workbook_catalog, split_list, split_steps

WORKBOOK = PROJECT_ROOT / "data" / "workflows.xlsx"
ORIGINAL = PROJECT_ROOT / "AI_Agent_Workflow_Assessment (1).xlsx"
#: SHA-256 of the workbook exactly as received with the assignment.
ORIGINAL_SHA256 = "581be7682eeb4d3039caaaa06595d197c103f023d7ef02d59c481ad31110f6ce"


def make_book(path: Path, sheets: dict[str, list[list]], merges: dict[str, list[str]] | None = None) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        for rng in (merges or {}).get(name, []):
            ws.merge_cells(rng)
    wb.save(path)
    return path


HEADER = ["Workflow_ID", "Workflow_Name", "Trigger", "Inputs", "Steps", "Decision_Logic", "Tools_Required", "Expected_Output"]


# ------------------------------------------------------------- the real file
def test_workbook_is_the_unmodified_assignment_file():
    assert hashlib.sha256(WORKBOOK.read_bytes()).hexdigest() == ORIGINAL_SHA256
    if ORIGINAL.exists():
        assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest() == ORIGINAL_SHA256


def test_all_ten_workflows_are_extracted():
    cat = load_workbook_catalog(WORKBOOK)
    assert [w.workflow_id for w in cat.workflows] == list(REQUIRED_WORKFLOW_IDS)
    assert len({w.workflow_id for w in cat.workflows}) == 10
    assert cat.warnings == []
    assert all(w.name and w.trigger and w.steps and w.inputs and w.tools_required and w.expected_output and w.decision_logic
               for w in cat.workflows)


def test_every_sheet_is_classified_and_every_row_read():
    cat = load_workbook_catalog(WORKBOOK)
    roles = {s.name: (s.role, s.rows_read) for s in cat.sheets}
    assert roles == {"Workflows": ("workflows", 10), "Test_Questions": ("test_questions", 10)}
    assert len(cat.test_questions) == 10
    assert {t.workflow_id for t in cat.test_questions} == set(REQUIRED_WORKFLOW_IDS)


def test_step_counts_and_splitting_match_the_workbook():
    cat = load_workbook_catalog(WORKBOOK)
    counts = {w.workflow_id: len(w.steps) for w in cat.workflows}
    assert counts == {"WF001": 5, "WF002": 5, "WF003": 6, "WF004": 5, "WF005": 5,
                      "WF006": 6, "WF007": 6, "WF008": 6, "WF009": 6, "WF010": 6}
    wf1 = cat.workflows[0]
    assert wf1.steps[0] == "Load inventory"
    assert wf1.inputs == ["Product inventory CSV", "minimum stock threshold"]
    assert wf1.tools_required == ["CSV reader", "calculator"]
    assert "current_stock < minimum_stock" in wf1.decision_logic
    assert (wf1.sheet, wf1.row) == ("Workflows", 2)


def test_test_questions_text_is_preserved():
    cat = load_workbook_catalog(WORKBOOK)
    q = {t.workflow_id: t.request for t in cat.test_questions}
    assert q["WF005"] == "Where is order ORD-1001?"
    assert q["WF010"] == "Which workflows are failing most often?"


# ----------------------------------------------------------- splitting helpers
@pytest.mark.parametrize("text,expected", [
    ("a → b → c", ["a", "b", "c"]),
    ("a -> b -> c", ["a", "b", "c"]),
    ("a\nb\nc", ["a", "b", "c"]),
    ("1. first\n2. second\n3) third", ["first", "second", "third"]),
    ("- one\n- two", ["one", "two"]),
    ("a; b; c", ["a", "b", "c"]),
    ("single step", ["single step"]),
    ("", []),
    ("a →  → b", ["a", "b"]),
])
def test_split_steps(text, expected):
    assert split_steps(text) == expected


def test_split_list_uses_semicolons_not_commas():
    assert split_list("Order ID or customer email") == ["Order ID or customer email"]
    assert split_list("Campaign goal; product list; dates") == ["Campaign goal", "product list", "dates"]
    assert split_list("x, y; z") == ["x, y", "z"]


# ------------------------------------------------------------ synthetic layouts
def test_headers_may_be_reordered_renamed_and_not_in_row_one(tmp_path):
    rows = [
        ["Company workflow catalogue"], [],
        ["Process Steps", "Workflow ID", "Name", "Rules", "Tools", "Output", "Input", "Trigger"],
        ["A → B", "X1", "Demo", "if x", "calc", "report", "file", "user asks"],
    ]
    cat = load_workbook_catalog(make_book(tmp_path / "b.xlsx", {"Sheet1": rows}))
    wf = cat.workflows[0]
    assert (wf.workflow_id, wf.name, wf.steps, wf.decision_logic, wf.tools_required, wf.expected_output, wf.inputs, wf.trigger) == \
        ("X1", "Demo", ["A", "B"], "if x", ["calc"], "report", ["file"], "user asks")
    assert cat.sheets[0].header_row == 3


def test_merged_cells_are_expanded(tmp_path):
    rows = [HEADER,
            ["W1", "Name", "T", "i", "s1 → s2", "d", "t", "o"],
            ["W2", "Other", "T", "i", "s1", "d", "t", "o"]]
    # merge the Trigger cell of both rows: the value must apply to both workflows
    cat = load_workbook_catalog(make_book(tmp_path / "m.xlsx", {"Workflows": rows}, {"Workflows": ["C2:C3"]}))
    assert [w.trigger for w in cat.workflows] == ["T", "T"]


def test_rows_without_an_id_continue_the_previous_workflow(tmp_path):
    rows = [HEADER, ["W1", "Name", "T", "i", "s1 → s2", "d", "t", "o"], [None, None, None, None, "s3", None, None, None]]
    cat = load_workbook_catalog(make_book(tmp_path / "c.xlsx", {"Workflows": rows}))
    assert len(cat.workflows) == 1
    assert cat.workflows[0].steps == ["s1", "s2", "s3"]
    assert any("continuation" in w for w in cat.warnings)


def test_unknown_columns_are_kept_not_dropped(tmp_path):
    rows = [HEADER + ["Owner"], ["W1", "Name", "T", "i", "s", "d", "t", "o", "ops team"]]
    cat = load_workbook_catalog(make_book(tmp_path / "e.xlsx", {"Workflows": rows}))
    assert cat.workflows[0].extra == {"Owner": "ops team"}
    assert cat.sheets[0].unrecognized_columns == ["Owner"]
    assert any("Owner" in w for w in cat.warnings)


def test_multiline_cells_and_numbered_steps(tmp_path):
    rows = [HEADER, ["W1", "Name", "T", "file one\nfile two", "1. read\n2. check\n3. write", "d", "t1;\nt2", "o"]]
    wf = load_workbook_catalog(make_book(tmp_path / "n.xlsx", {"Workflows": rows})).workflows[0]
    assert wf.steps == ["read", "check", "write"]
    assert wf.inputs == ["file one", "file two"]
    assert wf.tools_required == ["t1", "t2"]


def test_sheet_without_headers_is_ignored_with_a_warning(tmp_path):
    book = make_book(tmp_path / "w.xlsx", {"Workflows": [HEADER, ["W1", "N", "T", "i", "s", "d", "t", "o"]], "Notes": [["random", "text"]]})
    cat = load_workbook_catalog(book)
    assert {s.name: s.role for s in cat.sheets} == {"Workflows": "workflows", "Notes": "ignored"}
    assert any("Notes" in w for w in cat.warnings)


def test_incomplete_test_questions_are_reported(tmp_path):
    book = make_book(tmp_path / "t.xlsx", {
        "Workflows": [HEADER, ["W1", "N", "T", "i", "s", "d", "t", "o"]],
        "Test_Questions": [["Workflow_ID", "Test_Request", "What_To_Check"], ["W1", None, "x"], ["W1", "do it", "y"]]})
    cat = load_workbook_catalog(book)
    assert len(cat.test_questions) == 1
    assert any("incomplete test question" in w for w in cat.warnings)


# ------------------------------------------------------------------- failures
def test_missing_file_is_reported_clearly(tmp_path):
    with pytest.raises(CatalogError, match="not found") as exc:
        load_workbook_catalog(tmp_path / "nope.xlsx")
    assert exc.value.code == "workbook_missing"


def test_unsupported_extension(tmp_path):
    f = tmp_path / "x.csv"
    f.write_text("a,b")
    with pytest.raises(CatalogError, match="Unsupported workbook type"):
        load_workbook_catalog(f)


def test_corrupt_workbook(tmp_path):
    f = tmp_path / "bad.xlsx"
    f.write_text("this is not a zip file")
    with pytest.raises(CatalogError) as exc:
        load_workbook_catalog(f)
    assert exc.value.code == "workbook_unreadable"
    assert "bad.xlsx" in exc.value.user_message


def test_workbook_without_workflows(tmp_path):
    f = make_book(tmp_path / "e.xlsx", {"Sheet": [["foo", "bar"], [1, 2]]})
    with pytest.raises(CatalogError, match="No workflow definitions"):
        load_workbook_catalog(f)
