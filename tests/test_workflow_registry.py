"""Registry: loading, lookup, and static validation of definitions."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import openpyxl
import pytest

from src.config import REQUIRED_WORKFLOW_IDS
from src.errors import CatalogError
from src.workflow_registry import WorkflowRegistry
from tests.conftest import EXTENSION_DIR, edit_spec


def load(settings, tools) -> WorkflowRegistry:
    reg = WorkflowRegistry(settings, tools)
    reg.load()
    return reg


def errors(reg: WorkflowRegistry, wid: str | None = None) -> list[str]:
    return [i.message for i in reg.report.errors if wid is None or i.workflow_id == wid]


def edit_workbook(settings, fn):
    wb = openpyxl.load_workbook(settings.workbook_path)
    fn(wb["Workflows"])
    wb.save(settings.workbook_path)


# ------------------------------------------------------------------ happy path
def test_all_ten_required_workflows_register_cleanly(registry):
    rep = registry.report
    assert rep.ok and rep.issues == []
    assert rep.loaded_ids == list(REQUIRED_WORKFLOW_IDS)
    assert rep.missing_ids == [] and rep.invalid_ids == []
    assert rep.status_label.startswith("OK - 10/10")
    assert len(rep.workbook_sha256) == 64


def test_lookup_and_unknown_id(registry):
    d = registry.get("WF005")
    assert d.name == "Customer Order Status"
    assert registry.has("WF005") and not registry.has("WF999")
    with pytest.raises(CatalogError) as exc:
        registry.get("WF999")
    assert exc.value.code == "unknown_workflow"
    with pytest.raises(CatalogError):
        registry.get(None)  # type: ignore[arg-type]


def test_definition_keeps_the_excel_text_verbatim(registry):
    d = registry.get("WF002")
    assert d.decision_logic == "Flag when price difference exceeds 10%"
    assert d.success_output.startswith("Validation report showing matched products")
    assert d.excel_steps[0] == "Load product prices"
    assert d.tools_required == ["CSV reader", "calculator"]
    assert d.test_requests[0].request == "Find products where vendor price differs by more than 10%."
    assert "Product Price Validation" in d.name and d.source.endswith("Workflows row 3")


def test_router_catalog_exposes_only_what_the_router_needs(registry):
    cat = registry.router_catalog()
    assert [c["id"] for c in cat] == list(REQUIRED_WORKFLOW_IDS)
    wf3 = next(c for c in cat if c["id"] == "WF003")
    assert set(wf3) == {"id", "name", "description", "triggers", "keywords", "required_inputs", "inputs"}
    assert wf3["required_inputs"] == ["vendor_file"]


def test_every_workflow_defines_conditions_inputs_and_final_outcome(registry):
    for d in registry.all():
        assert d.conditions, d.id
        assert d.inputs and d.steps and "completed" in d.final.outcomes
        assert d.allowed_tools


def test_every_step_maps_to_an_excel_step_in_order(registry):
    for d in registry.all():
        excel = [s.lower().strip(" .") for s in d.excel_steps]
        seen: list[str] = []
        for s in d.steps:
            n = s.excel_step.lower().strip(" .")
            if not seen or seen[-1] != n:
                seen.append(n)
        assert seen == excel, d.id


def test_normalized_catalog_is_json_serialisable(registry):
    data = registry.normalized_catalog()
    assert len(data["workflows"]) == 10
    json.dumps(data)  # must not raise


def test_normalized_json_file_is_in_sync_with_the_workbook(registry):
    path = Path(__file__).resolve().parent.parent / "data" / "normalized_workflows.json"
    assert path.exists(), "run `python scripts/validate_workflows.py --write` to regenerate it"
    assert json.loads(path.read_text()) == json.loads(json.dumps(registry.normalized_catalog(), default=str))


# ------------------------------------------------------------------ reload
def test_reload_picks_up_changes_without_new_code(catalog_copy, tools):
    reg = load(catalog_copy, tools)
    assert reg.get("WF001").inputs[2].default == 2.0
    edit_spec(catalog_copy, "WF001", "default: 2.0", "default: 3.0")
    reg.reload()
    assert reg.get("WF001").inputs[2].default == 3.0 and reg.report.ok


# ------------------------------------------------------------- invalid catalogs
def test_duplicate_ids_in_the_workbook_are_rejected(catalog_copy, tools):
    edit_workbook(catalog_copy, lambda ws: ws.append([c.value for c in ws[4]]))  # copy of WF003
    reg = load(catalog_copy, tools)
    assert not reg.report.ok and "WF003" in reg.report.missing_ids
    assert any("Duplicate workflow ID WF003" in m for m in errors(reg))
    assert not reg.has("WF003")


def test_missing_spec_file(catalog_copy, tools):
    (catalog_copy.specs_dir / "WF004.yaml").unlink()
    reg = load(catalog_copy, tools)
    assert not reg.report.ok and reg.report.missing_ids == ["WF004"]
    assert any("no execution spec" in m for m in errors(reg, "WF004"))


def test_required_workflow_missing_from_the_workbook(catalog_copy, tools):
    edit_workbook(catalog_copy, lambda ws: ws.delete_rows(11))
    reg = load(catalog_copy, tools)
    assert not reg.report.ok and reg.report.missing_ids == ["WF010"]
    assert reg.has("WF009")  # the others still work
    assert reg.report.status_label.startswith("DEGRADED")


def test_workbook_step_changed_means_spec_is_stale(catalog_copy, tools):
    edit_workbook(catalog_copy, lambda ws: setattr(ws["E2"], "value", ws["E2"].value.replace("Load inventory", "Fetch stock")))
    reg = load(catalog_copy, tools)
    msgs = errors(reg, "WF001")
    assert any("Excel step 1" in m and "Fetch stock" in m for m in msgs)
    assert not reg.has("WF001")


def test_workbook_gains_a_step_the_spec_does_not_implement(catalog_copy, tools):
    edit_workbook(catalog_copy, lambda ws: setattr(ws["E2"], "value", ws["E2"].value + " → send an email"))
    assert any("Excel step 6" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_excel_input_not_covered_by_the_spec(catalog_copy, tools):
    edit_workbook(catalog_copy, lambda ws: setattr(ws["D2"], "value", ws["D2"].value + "; supplier email"))
    assert any("'supplier email' is not covered" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_spec_claims_to_cover_an_input_the_workbook_does_not_have(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", 'covers: ["minimum stock threshold"]', 'covers: ["minimum stock threshold", "ghost input"]')
    assert any("ghost input" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_excel_tool_without_mapping_and_extra_mapping(catalog_copy, tools):
    edit_workbook(catalog_copy, lambda ws: setattr(ws["G2"], "value", ws["G2"].value + "; OCR engine"))
    assert any("'OCR engine' is not mapped" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_spec_maps_to_a_tool_that_is_not_registered(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", '"CSV reader": [csv_reader]', '"CSV reader": [csv_reader_v2]')
    assert any("csv_reader_v2" in m and "not registered" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_step_uses_an_unregistered_tool(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", "tool: csv_reader", "tool: warp_drive")
    assert any("tool 'warp_drive' is not registered" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_step_uses_a_tool_the_workflow_may_not_use(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", "tool: csv_reader", "tool: order_search")
    assert any("not permitted for this workflow" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_unknown_and_missing_tool_parameters(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", "      source: \"=inputs.inventory_file\"\n", "      source: \"=inputs.inventory_file\"\n      colour: red\n")
    assert any("unknown parameter 'colour'" in m for m in errors(load(catalog_copy, tools), "WF001"))
    edit_spec(catalog_copy, "WF001", "      source: \"=inputs.inventory_file\"\n      colour: red\n", "")
    assert any("required parameter 'source'" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_invalid_expression_syntax_is_caught_at_load_time(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", 'where: "row.current_stock < row.threshold"', 'where: "row.current_stock <"')
    assert any("invalid expression" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_unknown_function_and_eval_attempt_are_rejected(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", 'where: "row.current_stock < row.threshold"', 'where: "__import__(\'os\')"')
    assert any("Unknown function" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_reference_to_unknown_step(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", 'rows: "=steps.compare_stock.rows"', 'rows: "=steps.nonexistent.rows"')
    assert any("unknown step 'nonexistent'" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_reference_to_a_step_that_runs_later(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", 'rows: "=steps.compare_stock.rows"', 'rows: "=steps.totals.rows"')
    assert any("does not run before it" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_reference_to_a_field_the_tool_does_not_output(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", 'rows: "=steps.compare_stock.rows"', 'rows: "=steps.compare_stock.records"')
    assert any("no output field 'records'" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_reference_to_an_undeclared_input(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", 'override: "=inputs.minimum_stock"', 'override: "=inputs.min_stock"')
    assert any("unknown input 'min_stock'" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_row_expression_cannot_reach_inputs_or_steps(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", 'where: "row.current_stock < row.threshold"', 'where: "row.current_stock < inputs.minimum_stock"')
    assert any("'inputs.minimum_stock' is not available here" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_vars_must_be_defined_in_the_step(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", "threshold: \"coalesce(vars.override, row.minimum_stock)\"", "threshold: \"coalesce(vars.nope, row.minimum_stock)\"")
    assert any("'vars.nope' is used but not defined" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_goto_backwards_is_rejected(catalog_copy, tools):
    edit_spec(catalog_copy, "WF005", "then: {type: continue, message: \"Recognised as {{ steps.validate_identifier.kind }}.\"}",
              "then: {type: goto, target: validate_identifier}")
    assert any("jumps backwards" in m for m in errors(load(catalog_copy, tools), "WF005"))


def test_end_action_needs_a_defined_outcome(catalog_copy, tools):
    edit_spec(catalog_copy, "WF009", "outcome: escalated", "outcome: panic")
    assert any("end outcome 'panic' is not defined" in m for m in errors(load(catalog_copy, tools), "WF009"))


def test_non_completed_outcome_may_not_use_steps_that_have_not_run(catalog_copy, tools):
    edit_spec(catalog_copy, "WF009", 'why_nobody_qualifies: "=steps.rank_candidates.rows"', 'why_nobody_qualifies: "=steps.pick_employee.rows"')
    assert any("has not run when this outcome is produced" in m for m in errors(load(catalog_copy, tools), "WF009"))


def test_final_result_reference_to_unknown_step(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", 'products_checked: "=steps.load_inventory.row_count"', 'products_checked: "=steps.ghost.row_count"')
    assert any("unknown step 'ghost'" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_duplicate_step_ids_and_invalid_yaml(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", "id: compare_stock", "id: load_inventory")
    assert any("used more than once" in m for m in errors(load(catalog_copy, tools), "WF001"))
    (catalog_copy.specs_dir / "WF002.yaml").write_text("workflow_id: [unclosed")
    assert any("not valid YAML" in m for m in errors(load(catalog_copy, tools), "WF002"))


def test_spec_with_unknown_keys_is_refused(catalog_copy, tools):
    edit_spec(catalog_copy, "WF001", "workflow_id: WF001", "workflow_id: WF001\nsurprise: true")
    assert any("WF001.yaml is invalid" in m for m in errors(load(catalog_copy, tools), "WF001"))


def test_workbook_not_found_means_an_unhealthy_catalog(tmp_path, settings, tools):
    reg = load(settings.with_overrides(workbook_path=tmp_path / "missing.xlsx"), tools)
    assert not reg.report.ok and len(reg.report.missing_ids) == 10
    assert any("not found" in m for m in errors(reg))
    assert reg.all() == []


def test_disabled_workflow_is_listed_but_not_offered_to_the_router(catalog_copy, tools):
    edit_spec(catalog_copy, "WF010", "workflow_id: WF010", "workflow_id: WF010\nenabled: false")
    reg = load(catalog_copy, tools)
    assert reg.report.ok and "WF010" not in reg.ids()
    assert reg.get("WF010").enabled is False and len(reg.all(include_disabled=True)) == 10


# --------------------------------------------------------------- extension source
def test_extension_workflow_registers_alongside_the_ten(settings, tools):
    reg = load(settings.with_overrides(extra_catalog_dirs=(EXTENSION_DIR,)), tools)
    assert reg.report.ok and reg.report.extension_ids == ["WF011"]
    assert reg.get("WF011").is_extension and not reg.get("WF001").is_extension
    assert reg.report.status_label.startswith("OK - 10/10")
    assert "WF011" not in reg.normalized_catalog()["workflows"].__repr__()  # the committed catalog stays the assignment's ten


def test_extension_cannot_override_an_assignment_workflow(tmp_path, settings, tools):
    ext = tmp_path / "ext"
    (ext / "specs").mkdir(parents=True)
    shutil.copy(settings.specs_dir / "WF001.yaml", ext / "specs" / "WF001.yaml")
    wb = openpyxl.load_workbook(settings.workbook_path)
    for name in list(wb.sheetnames):
        if name != "Workflows":
            del wb[name]
    ws = wb["Workflows"]
    ws.delete_rows(3, 20)  # keep header + WF001 only
    wb.save(ext / "dup.xlsx")
    reg = load(settings.with_overrides(extra_catalog_dirs=(ext,)), tools)
    assert any("Duplicate workflow ID WF001" in m for m in errors(reg, "WF001"))
    assert reg.get("WF001").source.startswith("workflows.xlsx")  # the original is untouched


def test_missing_extension_folder_is_only_a_warning(tmp_path, settings, tools):
    reg = load(settings.with_overrides(extra_catalog_dirs=(tmp_path / "nowhere",)), tools)
    assert reg.report.ok and any("no .xlsx" in i.message for i in reg.report.warnings)
