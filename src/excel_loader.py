"""Read the assignment workbook into :class:`ExcelCatalog`.

The loader does **not** assume a fixed layout. It

* scans every worksheet for a header row (within the first rows) and decides the
  sheet's role from the headers it finds (workflow definitions vs. test questions),
* matches headers by normalised name/aliases (``Workflow_ID`` == ``workflow id``),
* un-merges merged cells (the merged value is filled into every covered cell),
* treats a row without an ID that still has content as a continuation of the
  previous workflow (multi-row workflows),
* splits list-like cells (``;`` / line breaks) and step cells (``→`` / ``->`` /
  line breaks / numbered lists) into ordered items,
* keeps unrecognised columns instead of silently dropping them, and
* records the sheet + row each workflow came from.

Anything surprising becomes an entry in ``ExcelCatalog.warnings``.
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .errors import CatalogError
from .models import ExcelCatalog, ExcelTestQuestion, ExcelWorkflow, SheetInfo

HEADER_SCAN_ROWS = 15

# canonical field -> accepted normalised header names
WORKFLOW_COLUMNS: dict[str, set[str]] = {
    "workflow_id": {"workflowid", "id", "wfid", "workflowcode", "workflowno"},
    "name": {"workflowname", "name", "title", "workflowtitle", "processname"},
    "trigger": {"trigger", "triggers", "userrequest", "triggerevent"},
    "inputs": {"inputs", "input", "requiredinputs", "parameters", "requiredparameters"},
    "steps": {"steps", "step", "executionsteps", "process", "workflowsteps", "processsteps", "procedure", "stepsequence", "workflowprocess"},
    "decision_logic": {"decisionlogic", "decision", "decisions", "conditions", "rules", "businessrules"},
    "tools": {"toolsrequired", "tools", "tool", "apis", "toolsapis", "requiredtools"},
    "expected_output": {"expectedoutput", "output", "expectedresult", "result", "outputs"},
}
TEST_COLUMNS: dict[str, set[str]] = {
    "workflow_id": WORKFLOW_COLUMNS["workflow_id"],
    "test_request": {"testrequest", "request", "question", "testquestion", "testprompt", "prompt"},
    "what_to_check": {"whattocheck", "check", "checks", "expected", "notes", "purpose"},
}
MIN_WORKFLOW_COLUMNS = {"workflow_id", "name", "steps"}

_ARROW_SPLIT = re.compile(r"\s*(?:→|⟶|⇒|->|=>)\s*")
_BULLET = re.compile(r"^\s*(?:\d+[\.\)]|[-•*–])\s+")


def _norm_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    return "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")).strip()


def split_steps(text: str) -> list[str]:
    """Split a Steps cell into ordered step labels."""
    text = _clean(text)
    if not text:
        return []
    if _ARROW_SPLIT.search(text):
        # a line break always ends a step, so a continuation row cannot be glued onto the previous step
        parts = [p for line in text.split("\n") for p in _ARROW_SPLIT.split(line)]
    elif "\n" in text:
        parts = text.split("\n")
    elif ";" in text:
        parts = text.split(";")
    else:
        parts = [text]
    return [p for p in (_BULLET.sub("", part).strip(" .") for part in parts) if p]


def split_list(text: str) -> list[str]:
    """Split an Inputs / Tools cell into items (``;`` or line breaks)."""
    text = _clean(text)
    if not text:
        return []
    parts = re.split(r"[;\n]+", text)
    return [p for p in (_BULLET.sub("", part).strip(" .") for part in parts) if p]


def _sheet_grid(ws: Worksheet) -> list[list[Any]]:
    """Cell values with merged ranges expanded."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:
        top = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if r - 1 < len(rows) and c - 1 < len(rows[r - 1]):
                    rows[r - 1][c - 1] = top
    return rows


def _find_header(grid: list[list[Any]], spec: dict[str, set[str]], required: set[str]) -> tuple[int, dict[str, int]] | None:
    for idx, row in enumerate(grid[:HEADER_SCAN_ROWS]):
        found: dict[str, int] = {}
        for col, cell in enumerate(row):
            norm = _norm_header(cell)
            if not norm:
                continue
            for canonical, aliases in spec.items():
                if norm in aliases and canonical not in found:
                    found[canonical] = col
        if required <= set(found):
            return idx, found
    return None


def file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_workbook_catalog(path: str | Path) -> ExcelCatalog:
    """Parse ``path`` into an :class:`ExcelCatalog` (raises :class:`CatalogError`)."""
    path = Path(path)
    if not path.exists():
        raise CatalogError(
            f"Workflow workbook not found: {path}. Place the assignment workbook there (or set WORKFLOW_WORKBOOK).",
            code="workbook_missing",
        )
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise CatalogError(f"Unsupported workbook type '{path.suffix}'. Expected .xlsx", code="workbook_unsupported")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:  # corrupt / password-protected files
        raise CatalogError(f"Could not open workbook '{path.name}': {type(exc).__name__}", detail=str(exc), code="workbook_unreadable") from exc

    catalog = ExcelCatalog(path=str(path), sha256=file_sha256(path))
    for ws in wb.worksheets:
        grid = _sheet_grid(ws)
        info = SheetInfo(name=ws.title, role="ignored", dimensions=ws.dimensions or "")
        if ws.sheet_state != "visible":
            info.note = f"sheet is {ws.sheet_state}; read anyway"
        header = _find_header(grid, WORKFLOW_COLUMNS, MIN_WORKFLOW_COLUMNS)
        if header:
            info.role = "workflows"
            _read_workflow_rows(grid, header, ws.title, info, catalog)
        else:
            header = _find_header(grid, TEST_COLUMNS, {"workflow_id", "test_request"})
            if header:
                info.role = "test_questions"
                _read_test_rows(grid, header, ws.title, info, catalog)
            else:
                info.note = (info.note + "; " if info.note else "") + "no recognisable header row - sheet ignored"
                catalog.warnings.append(f"Sheet '{ws.title}' ignored: no recognisable workflow/test header row.")
        catalog.sheets.append(info)

    if not catalog.workflows:
        seen = "; ".join(f"'{s.name}' ({s.note or 'no header'})" for s in catalog.sheets) or "no sheets"
        raise CatalogError(
            f"No workflow definitions found in '{path.name}'. Expected a sheet whose header row has columns such as "
            f"Workflow_ID, Workflow_Name and Steps. Sheets examined: {seen}.",
            code="no_workflows",
        )
    return catalog


def _cell(row: list[Any], col: int | None) -> str:
    if col is None or col >= len(row):
        return ""
    return _clean(row[col])


def _read_workflow_rows(grid, header, sheet: str, info: SheetInfo, catalog: ExcelCatalog) -> None:
    hdr_idx, cols = header
    info.header_row = hdr_idx + 1
    header_cells = grid[hdr_idx]
    recognised_cols = set(cols.values())
    info.columns = [_clean(c) for c in header_cells if _clean(c)]
    extra_cols = {i: _clean(c) for i, c in enumerate(header_cells) if _clean(c) and i not in recognised_cols}
    info.unrecognized_columns = list(extra_cols.values())
    for missing in sorted(set(WORKFLOW_COLUMNS) - set(cols)):
        catalog.warnings.append(f"Sheet '{sheet}': column for '{missing}' not found; that field will be empty.")
    if extra_cols:
        catalog.warnings.append(f"Sheet '{sheet}': extra columns kept as notes: {', '.join(extra_cols.values())}.")

    current: dict[str, Any] | None = None

    def flush():
        nonlocal current
        if current is None:
            return
        wf = _build_workflow(current, sheet)
        catalog.workflows.append(wf)
        info.rows_read += 1
        current = None

    for offset, row in enumerate(grid[hdr_idx + 1 :], start=hdr_idx + 2):
        values = {key: _cell(row, col) for key, col in cols.items()}
        extra = {name: _cell(row, i) for i, name in extra_cols.items() if _cell(row, i)}
        if not any(values.values()) and not extra:
            continue  # blank row
        if values.get("workflow_id"):
            flush()
            current = {**values, "extra": extra, "row": offset}
        elif current is not None:
            # continuation of the previous workflow (multi-row definition)
            for key, val in values.items():
                if val:
                    current[key] = (current.get(key, "") + "\n" + val).strip()
            current["extra"].update(extra)
            catalog.warnings.append(
                f"Sheet '{sheet}' row {offset}: no Workflow_ID; treated as continuation of {current['workflow_id']}."
            )
        else:
            catalog.warnings.append(f"Sheet '{sheet}' row {offset}: content without a Workflow_ID before the first workflow; skipped.")
    flush()


def _build_workflow(raw: dict[str, Any], sheet: str) -> ExcelWorkflow:
    return ExcelWorkflow(
        workflow_id=raw["workflow_id"].strip(),
        name=raw.get("name", "").strip(),
        trigger=raw.get("trigger", ""),
        inputs_text=raw.get("inputs", ""),
        inputs=split_list(raw.get("inputs", "")),
        steps_text=raw.get("steps", ""),
        steps=split_steps(raw.get("steps", "")),
        decision_logic=raw.get("decision_logic", ""),
        tools_text=raw.get("tools", ""),
        tools_required=split_list(raw.get("tools", "")),
        expected_output=raw.get("expected_output", ""),
        sheet=sheet,
        row=raw["row"],
        extra=raw.get("extra", {}),
    )


def _read_test_rows(grid, header, sheet: str, info: SheetInfo, catalog: ExcelCatalog) -> None:
    hdr_idx, cols = header
    info.header_row = hdr_idx + 1
    info.columns = [_clean(c) for c in grid[hdr_idx] if _clean(c)]
    for offset, row in enumerate(grid[hdr_idx + 1 :], start=hdr_idx + 2):
        wf_id, req = _cell(row, cols.get("workflow_id")), _cell(row, cols.get("test_request"))
        if not wf_id and not req:
            continue
        if not wf_id or not req:
            catalog.warnings.append(f"Sheet '{sheet}' row {offset}: incomplete test question skipped.")
            continue
        catalog.test_questions.append(
            ExcelTestQuestion(
                workflow_id=wf_id,
                request=req,
                what_to_check=_cell(row, cols.get("what_to_check")),
                sheet=sheet,
                row=offset,
            )
        )
        info.rows_read += 1
